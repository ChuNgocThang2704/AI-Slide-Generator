import asyncio
import io
import json
import unittest

from services.lecture_script import (
    _drop_template_pictures, describe_figures, read_pptx, shrink_image,
    translate_scripts,
    NOTE_PRESENTER, SCENE_INTRO, SCENE_OUTRO, assemble_rows, estimated_minutes, has_cover, tidy_script, write_script,
)
from services.lecture_script_xlsx import build_workbook, build_workbook_of, sheet_name

LONG = "Mời các em quan sát nội dung này và ghi nhớ ý chính của phần học hôm nay nhé."


def _slides(count=4, cover=True):
    slides = [{"number": 1, "title": "Chương 2", "text": "Chương 2\nCác dạng tấn công" if cover else "từ " * 80, "notes": "", "has_figure": False}]
    for number in range(2, count + 1):
        slides.append({
            "number": number, "title": f"Mục {number}", "text": f"Nội dung tiếng Việt của mục {number} về mối đe dọa và lỗ hổng",
            "notes": "", "has_figure": number == 3,
        })
    return slides


class _Extractor:
    """Answers the slide batches and the intro/outro call; `skip` leaves those slides out."""

    def __init__(self, skip=()):
        self.skip = set(skip)
        self.calls = []

    async def _llm_completion_plain_text(self, messages, **kwargs):
        self.calls.append(kwargs.get("call_purpose"))
        if kwargs.get("call_purpose") == "lecture_script_frame":
            self.frame_prompt = messages[0]["content"]
            self.frame_payload = json.loads(messages[1]["content"])
        payload = json.loads(messages[1]["content"])
        if kwargs.get("call_purpose") == "lecture_script_frame":
            return json.dumps({"title": "Mối đe dọa và Lỗ hổng", "intro": "Chào các em.\n\n" + LONG, "outro": LONG})
        return json.dumps({"items": [
            {"slide": slide["slide"], "script": f"**Slide {slide['slide']}** mới.\n- {LONG}"}
            for slide in payload["slides"] if slide["slide"] not in self.skip
        ]})


class LectureScriptTests(unittest.TestCase):
    def test_spoken_text_loses_markdown_bullets_and_stage_directions(self):
        self.assertEqual(tidy_script("**Chào** các em.\n- ý một\n\n\n[cười] Hai."), "Chào các em.\n\ný một\n\nHai.")

    def test_title_slide_becomes_the_opening_scene(self):
        slides = _slides()
        self.assertTrue(has_cover(slides))
        rows = assemble_rows(slides, {2: "a", 3: "b", 4: "c"}, {"intro": "i", "outro": "o"})
        self.assertEqual([row["scene"] for row in rows], [SCENE_INTRO, "Slide 2", "Slide 3", "Slide 4", SCENE_OUTRO])
        self.assertEqual([row["note"] for row in rows], ["", "", "", "", ""])  # left for the video team

    def test_a_first_slide_with_real_content_keeps_its_own_scene(self):
        slides = _slides(cover=False)
        self.assertFalse(has_cover(slides))
        rows = assemble_rows(slides, {}, {})
        self.assertEqual([row["scene"] for row in rows][:2], [SCENE_INTRO, "Slide 1"])

    def test_duration_follows_the_speaking_pace(self):
        self.assertEqual(estimated_minutes([{"script": "từ " * 1120}]), 8)
        self.assertEqual(estimated_minutes([{"script": ""}]), 0)

    def test_script_is_written_for_every_slide(self):
        script = asyncio.run(write_script(_Extractor(), _slides(8), filename="bai.pptx"))
        self.assertEqual(script["title"], "Mối đe dọa và Lỗ hổng")
        self.assertEqual(script["language"], "vi")
        self.assertEqual(len(script["rows"]), 9)
        self.assertEqual(script["missing"], [])
        self.assertTrue(script["rows"][1]["script"].startswith("Slide 2 mới."))
        self.assertNotIn("**", script["rows"][1]["script"])

    def test_only_the_first_video_introduces_the_lecturer_and_the_course(self):
        later = _Extractor()
        asyncio.run(write_script(later, _slides(4)))
        self.assertIn("không giới thiệu tên giảng viên", later.frame_prompt)
        self.assertNotIn("cover_slide_text", later.frame_payload)
        first = _Extractor()
        asyncio.run(write_script(first, _slides(4), first_video=True))
        self.assertIn("giới thiệu giảng viên và học phần", first.frame_prompt)
        self.assertIn("cover_slide_text", first.frame_payload)

    def test_a_slide_the_model_skips_is_reported_and_left_for_the_user(self):
        script = asyncio.run(write_script(_Extractor(skip={3}), _slides(5)))
        self.assertEqual(script["missing"], ["Slide 3"])
        self.assertEqual(script["rows"][2]["script"], "")

    def test_rewriting_keeps_the_old_text_of_a_slide_the_model_skips(self):
        previous = [
            {"scene": SCENE_INTRO, "slide": 1, "script": "Lời cũ mở đầu"},
            {"scene": "Slide 2", "slide": 2, "script": "Lời cũ slide hai"},
            {"scene": "Slide 3", "slide": 3, "script": "Lời cũ slide ba"},
            {"scene": SCENE_OUTRO, "slide": None, "script": "Lời cũ kết"},
        ]
        script = asyncio.run(write_script(_Extractor(skip={3}), _slides(3), prompt="ngắn hơn", previous_rows=previous))
        by_scene = {row["scene"]: row["script"] for row in script["rows"]}
        self.assertTrue(by_scene["Slide 2"].startswith("Slide 2 mới."))
        self.assertEqual(by_scene["Slide 3"], "Lời cũ slide ba")
        self.assertEqual(script["missing"], [])


    def test_a_slide_without_words_gets_no_invented_script(self):
        slides = _slides(5)
        slides[3]["text"] = ""
        extractor = _Extractor()
        script = asyncio.run(write_script(extractor, slides))
        self.assertEqual(script["missing"], ["Slide 4"])
        self.assertEqual(script["rows"][3]["script"], "")


def _png(width=1600, height=1000, color=(200, 30, 30)):
    from PIL import Image

    buffer = io.BytesIO()
    Image.new("RGB", (width, height), color).save(buffer, format="PNG")
    return buffer.getvalue()


class _Vision:
    """Answers the figure calls; `answers` maps a slide title to the reply; records what it was sent."""

    model_name = "gateway-model"

    def __init__(self, answers):
        self.answers = answers
        self.sent = []

    async def _llm_completion_plain_text(self, messages, **kwargs):
        content = messages[0]["content"]
        self.sent.append({"kwargs": kwargs, "images": [part for part in content if part["type"] == "image_url"], "text": content[0]["text"]})
        for title, reply in self.answers.items():
            if f'"{title}"' in content[0]["text"]:
                if isinstance(reply, Exception):
                    raise reply
                return reply
        return "TRANG TRÍ"


class FigureTests(unittest.TestCase):
    def test_big_pictures_are_read_from_a_pptx_and_icons_and_backgrounds_are_not(self):
        from pptx import Presentation
        from pptx.util import Emu

        deck = Presentation()
        deck.slide_width, deck.slide_height = Emu(9144000), Emu(5143500)
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        png = _png(400, 300)
        slide.shapes.add_picture(io.BytesIO(png), Emu(500000), Emu(500000), Emu(4500000), Emu(3000000))   # a figure
        slide.shapes.add_picture(io.BytesIO(png), Emu(100000), Emu(100000), Emu(400000), Emu(400000))     # an icon
        slide.shapes.add_picture(io.BytesIO(png), 0, 0, deck.slide_width, deck.slide_height)               # a background
        out = io.BytesIO()
        deck.save(out)
        slides = read_pptx(out.getvalue())
        self.assertEqual(len(slides[0]["images"]), 1)
        self.assertTrue(slides[0]["has_figure"])

    def test_a_picture_is_shrunk_before_the_model_sees_it(self):
        import base64
        from PIL import Image

        url = shrink_image(_png(3000, 2000))
        self.assertTrue(url.startswith("data:image/jpeg;base64,"))
        shrunk = Image.open(io.BytesIO(base64.b64decode(url.split(",", 1)[1])))
        self.assertLessEqual(max(shrunk.size), 768)
        self.assertIsNone(shrink_image(b"not an image"))

    def test_described_figures_reach_the_slide_and_the_pictures_are_dropped(self):
        slides = _slides(4)
        for index, slide in enumerate(slides[1:]):
            slide["images"] = [_png(color=(40 * index, 30, 30))]      # each its own picture, none a repeated logo
        vision = _Vision({"Mục 2": "Sơ đồ ba khối nối nhau bằng mũi tên: mối đe dọa khai thác lỗ hổng dẫn tới tấn công.", "Mục 3": "TRANG TRÍ", "Mục 4": RuntimeError("model down")})
        described = asyncio.run(describe_figures(vision, slides, True))
        self.assertEqual(described, 1)
        self.assertIn("mối đe dọa khai thác lỗ hổng", slides[1]["figure"])
        self.assertNotIn("figure", slides[2])        # a decoration
        self.assertNotIn("figure", slides[3])        # a failed call: written from the text alone
        self.assertTrue(all("images" not in slide for slide in slides))
        self.assertEqual(vision.sent[0]["kwargs"]["model_override"], "gateway-model")   # never the self-hosted primary host
        self.assertTrue(vision.sent[0]["images"][0]["image_url"]["url"].startswith("data:image/jpeg;base64,"))

    def test_a_powerpoint_chart_is_read_as_its_data(self):
        from pptx import Presentation
        from pptx.chart.data import CategoryChartData
        from pptx.enum.chart import XL_CHART_TYPE
        from pptx.util import Emu

        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        data = CategoryChartData()
        data.categories = ["2020", "2021", "2022"]
        data.add_series("Sản lượng", (1.76, 1.74, 1.85))
        slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Emu(500000), Emu(500000), Emu(5000000), Emu(3000000), data)
        out = io.BytesIO()
        deck.save(out)
        text = read_pptx(out.getvalue())[0]["text"]
        self.assertIn("Biểu đồ", text)
        self.assertIn("Sản lượng", text)
        self.assertIn("2021 = 1.74", text)

    def test_a_picture_repeated_on_several_slides_is_template_not_lesson(self):
        logo, figure = _png(color=(0, 0, 200)), _png(color=(0, 200, 0))
        slides = [{"images": [logo]}, {"images": [logo, figure]}, {"images": [logo]}, {"images": [figure]}]
        removed = _drop_template_pictures(slides)
        self.assertEqual(removed, 3)
        self.assertEqual([len(slide["images"]) for slide in slides], [0, 1, 0, 1])    # the figure is on two slides only

    def test_figures_can_be_switched_off(self):
        slides = _slides(3)
        slides[1]["images"] = [_png()]
        vision = _Extractor()
        script = asyncio.run(write_script(vision, slides, look_at_figures=False))
        self.assertEqual(script["figures_described"], 0)
        self.assertNotIn("images", slides[1])

    def test_the_title_slide_is_not_looked_at(self):
        slides = _slides(3)
        for index, slide in enumerate(slides):
            slide["images"] = [_png(color=(40 * index, 30, 30))]
        vision = _Vision({})
        asyncio.run(describe_figures(vision, slides, True))
        self.assertEqual(len(vision.sent), 2)

    def test_the_description_is_given_to_the_script_writer(self):
        slides = _slides(3)
        slides[1]["images"] = [_png()]
        seen = {}

        class Both(_Extractor, _Vision):
            def __init__(self):
                _Extractor.__init__(self)
                _Vision.__init__(self, {"Mục 2": "Sơ đồ ba khối nối bằng mũi tên, đủ dài để hợp lệ."})

            async def _llm_completion_plain_text(self, messages, **kwargs):
                if isinstance(messages[0]["content"], list):
                    return await _Vision._llm_completion_plain_text(self, messages, **kwargs)
                if kwargs.get("call_purpose") == "lecture_script_slides":
                    seen["payload"] = json.loads(messages[1]["content"])
                return await _Extractor._llm_completion_plain_text(self, messages, **kwargs)

        script = asyncio.run(write_script(Both(), slides))
        self.assertEqual(script["figures_described"], 1)
        first = seen["payload"]["slides"][0]
        self.assertEqual(first["figure_description"], "Sơ đồ ba khối nối bằng mũi tên, đủ dài để hợp lệ.")
        self.assertNotIn("figure_description", seen["payload"]["slides"][1])
        self.assertEqual(slides[1]["figure"], first["figure_description"])      # kept on the slide, for rewriting later


class _Translator:
    """Answers translate_strings: 'EN:' in front of each item; `drop` makes it return one item too few."""

    def __init__(self, drop=False):
        self.drop = drop
        self.calls = 0

    async def _llm_completion_plain_text(self, messages, **kwargs):
        self.calls += 1
        items = json.loads(messages[1]["content"])["items"]
        if self.drop and len(items) > 1:
            items = items[:-1]
        return json.dumps({"items": [f"EN:{item}" for item in items]})


class EnglishSubtitleTests(unittest.TestCase):
    def test_every_paragraph_is_translated_and_the_blank_line_structure_is_kept(self):
        result = asyncio.run(translate_scripts(_Translator(), ["Chào các em.\n\nHôm nay học bài mới.", "", "Một đoạn."]))
        self.assertEqual(result["failed"], [])
        self.assertEqual(result["items"], [
            {"i": 0, "en": "EN:Chào các em.\n\nEN:Hôm nay học bài mới."},
            {"i": 2, "en": "EN:Một đoạn."},
        ])

    def test_a_batch_the_model_gets_wrong_is_split_until_it_is_right(self):
        translator = _Translator(drop=True)
        result = asyncio.run(translate_scripts(translator, ["a\n\nb\n\nc\n\nd"]))
        self.assertEqual(result["failed"], [])
        self.assertEqual(result["items"][0]["en"], "EN:a\n\nEN:b\n\nEN:c\n\nEN:d")
        self.assertGreater(translator.calls, 2)

    def test_nothing_to_translate_gives_nothing(self):
        self.assertEqual(asyncio.run(translate_scripts(_Translator(), ["", "  "])), {"items": [], "failed": []})


class LectureScriptWorkbookTests(unittest.TestCase):
    def test_english_goes_in_a_fourth_column_only_when_there_is_some(self):
        try:
            from openpyxl import load_workbook
        except ImportError:
            self.skipTest("openpyxl is not installed")
        plain = load_workbook(io.BytesIO(build_workbook({"title": "T", "rows": [{"scene": "Slide 2", "script": "Một.", "note": ""}]}))).active
        self.assertIsNone(plain["D4"].value)
        self.assertEqual({str(c) for c in plain.merged_cells.ranges}, {"A1:C1", "A2:C2"})
        both = load_workbook(io.BytesIO(build_workbook({"title": "T", "rows": [
            {"scene": "Slide 2", "script": "Một.", "note": "", "en": "One."},
            {"scene": "Slide 3", "script": "Hai.", "note": ""},
        ]}))).active
        self.assertEqual([both[f"{c}4"].value for c in "ABCD"], ["PHÂN CẢNH", "LỜI THOẠI", "LỜI THOẠI (ENGLISH)", "LƯU Ý DỰNG"])
        self.assertEqual(both["B5"].value, "Một.")
        self.assertEqual(both["C5"].value, "One.")
        self.assertIsNone(both["C6"].value)
        self.assertEqual(both["D5"].value, None)
        self.assertEqual({str(c) for c in both.merged_cells.ranges}, {"A1:D1", "A2:D2"})

    def test_workbook_has_the_layout_of_a_production_script(self):
        try:
            from openpyxl import load_workbook
        except ImportError:
            self.skipTest("openpyxl is not installed")
        data = build_workbook({
            "title": "Video 4: Mối đe dọa / Lỗ hổng", "duration_minutes": 8,
            "rows": [
                {"scene": SCENE_INTRO, "script": "Chào các em.\n\nHôm nay…", "note": NOTE_PRESENTER},
                {"scene": "Slide 2", "script": "Nội dung.", "note": ""},
            ],
        })
        sheet = load_workbook(io.BytesIO(data)).active
        self.assertEqual(sheet["A1"].value, "Video 4: Mối đe dọa / Lỗ hổng")
        self.assertEqual(sheet["A2"].value, "Thời lượng dự kiến: 8 phút")
        self.assertEqual([sheet[f"{column}4"].value for column in "ABC"], ["PHÂN CẢNH", "LỜI THOẠI", "LƯU Ý DỰNG"])
        self.assertEqual(sheet["A5"].value, SCENE_INTRO)
        self.assertEqual(sheet["B5"].value, "Chào các em.\n\nHôm nay…")
        self.assertEqual(sheet["C5"].value, NOTE_PRESENTER)
        self.assertIsNone(sheet["C6"].value)
        self.assertEqual(sheet["B5"].font.name, "Times New Roman")
        self.assertTrue(sheet["B5"].alignment.wrap_text)
        self.assertEqual({str(cells) for cells in sheet.merged_cells.ranges}, {"A1:C1", "A2:C2"})

    def test_several_scripts_share_one_workbook_with_a_sheet_each(self):
        try:
            from openpyxl import load_workbook
        except ImportError:
            self.skipTest("openpyxl is not installed")
        row = [{"scene": "Slide 2", "script": "Nội dung.", "note": ""}]
        data = build_workbook_of([
            {"title": "Video 4: A", "sheet": "Video 4", "duration_minutes": 8, "rows": row},
            {"title": "Video 5: B", "sheet": "Video 5", "duration_minutes": 7, "rows": row},
            {"title": "Trùng tên", "sheet": "video 4", "duration_minutes": 1, "rows": row},
        ])
        workbook = load_workbook(io.BytesIO(data))
        self.assertEqual(workbook.sheetnames, ["Video 4", "Video 5", "video 4 (2)"])
        self.assertEqual(workbook["Video 5"]["A1"].value, "Video 5: B")
        self.assertEqual(workbook["Video 5"]["A2"].value, "Thời lượng dự kiến: 7 phút")

    def test_sheet_name_is_legal_in_excel(self):
        self.assertEqual(sheet_name("Video 4: A/B [x]"), "Video 4  A B  x")
        self.assertLessEqual(len(sheet_name("a" * 80)), 31)


if __name__ == "__main__":
    unittest.main()
