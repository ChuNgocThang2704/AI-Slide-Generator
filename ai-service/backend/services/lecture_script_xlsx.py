"""The lecture script as the Excel sheet a video team works from.

Layout of the reference files: a merged title row, a merged "Thời lượng dự kiến" row, then a
three-column table PHÂN CẢNH / LỜI THOẠI / LƯU Ý DỰNG in Times New Roman with a navy header.
"""

from __future__ import annotations

import io
import math
import re
from typing import Any, Dict, List

_FONT = "Times New Roman"
_NAVY = "FF003366"
_TITLE_BLUE = "FF000080"
_WIDTHS = {"A": 25, "B": 80, "C": 35, "D": 80}
_CHARS_PER_LINE = 92          # what an 80-wide column of 11pt Times New Roman holds
_LINE_POINTS = 15.0


def _row_height(script: str) -> float:
    lines = 0
    for paragraph in str(script or "").split("\n"):
        lines += max(1, math.ceil(len(paragraph) / _CHARS_PER_LINE))
    return min(409.0, max(30.0, lines * _LINE_POINTS + 6))


def sheet_name(title: str) -> str:
    name = re.sub(r"[\[\]:*?/\\]", " ", str(title or "")).strip()
    return (name or "Kịch bản")[:31]


def build_workbook(script: Dict[str, Any]) -> bytes:
    """.xlsx bytes for {"title", "sheet", "duration_minutes", "rows": [{"scene","script","note"}]}."""
    return build_workbook_of([script])


def build_workbook_of(scripts: List[Dict[str, Any]]) -> bytes:
    """One workbook with a sheet per script, the way a chapter's videos are kept together."""
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.remove(workbook.active)
    used: set = set()
    for script in scripts:
        base = sheet_name(script.get("sheet") or script.get("title") or "Kịch bản")
        name, counter = base, 2
        while name.lower() in used:     # Excel refuses two sheets with one name
            suffix = f" ({counter})"
            name, counter = base[: 31 - len(suffix)] + suffix, counter + 1
        used.add(name.lower())
        _fill_sheet(workbook.create_sheet(title=name), script)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _fill_sheet(sheet, script: Dict[str, Any]) -> None:
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    rows: List[Dict[str, Any]] = [row for row in (script.get("rows") or []) if isinstance(row, dict)]
    # English lines for the subtitles go in a fourth column; the first three stay exactly as in the template.
    with_english = any(str(row.get("en") or "").strip() for row in rows)
    columns = "ABCD" if with_english else "ABC"
    last = columns[-1]
    for column in columns:
        sheet.column_dimensions[column].width = _WIDTHS[column]

    thin = Side(style="thin", color="FF000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    sheet.merge_cells(f"A1:{last}1")
    sheet["A1"] = str(script.get("title") or "Kịch bản bài giảng")
    sheet["A1"].font = Font(name=_FONT, size=14, bold=True, color=_TITLE_BLUE)
    sheet["A1"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 30

    minutes = script.get("duration_minutes")
    sheet.merge_cells(f"A2:{last}2")
    sheet["A2"] = f"Thời lượng dự kiến: {minutes} phút" if minutes else "Thời lượng dự kiến:  phút"
    sheet["A2"].font = Font(name=_FONT, size=12, bold=True, color="FF333333")
    sheet["A2"].alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[2].height = 24.75
    sheet.row_dimensions[3].height = 9.75

    for column, heading in zip(columns, ("PHÂN CẢNH", "LỜI THOẠI", "LƯU Ý DỰNG", "LỜI THOẠI (ENGLISH)")):
        cell = sheet[f"{column}4"]
        cell.value = heading
        cell.font = Font(name=_FONT, size=11, bold=True, color="FFFFFFFF")
        cell.fill = PatternFill(fill_type="solid", fgColor=_NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    sheet.row_dimensions[4].height = 22

    for offset, row in enumerate(rows):
        index = 5 + offset
        values = (str(row.get("scene") or ""), str(row.get("script") or ""), str(row.get("note") or ""))
        if with_english:
            values += (str(row.get("en") or ""),)
        for column, value in zip(columns, values):
            cell = sheet[f"{column}{index}"]
            cell.value = value or None
            cell.font = Font(name=_FONT, size=11, bold=(column == "A"))
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = border
        sheet.row_dimensions[index].height = max(_row_height(values[1]), _row_height(values[3]) if with_english else 0)
