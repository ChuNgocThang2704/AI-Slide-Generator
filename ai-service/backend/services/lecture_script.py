"""Lecture script ("kịch bản dựng") for an uploaded slide deck.

The user brings a finished deck (PPTX or PDF). For each slide the lecturer's spoken words are
written, in the form a video team records from: an opening on camera, one scene per slide, a
closing on camera. Nothing here touches slide generation: the deck is only read.
"""

from __future__ import annotations

import asyncio
import io
import json
import re
from typing import Any, Awaitable, Callable, Dict, List, Optional

from services.content.json_utils import parse_json_response

MAX_SLIDES = 80
WORDS_PER_MINUTE = 140          # the pace of the reference scripts (about 1100 words for 8 minutes)
_BATCH_SLIDES = 5
_PARALLEL_BATCHES = 3
_SLIDE_TEXT_CHARS = 1400
_COVER_WORDS = 40               # a first slide with no more words than this is the title slide
_FIGURE_AREA = 0.15             # a picture this large (share of the slide) is a figure, not a logo or an icon
_BACKGROUND_AREA = 0.85         # ... and one this large is the slide's background, not something to talk about
_MAX_IMAGES_PER_SLIDE = 2       # the largest pictures of a slide are looked at
_MAX_FIGURE_SLIDES = 40         # slides of one deck whose figures are looked at
_FIGURE_SIDE = 768              # pictures are shrunk to this many pixels before the model sees them (cost)
_PARALLEL_FIGURES = 3
_TEMPLATE_REPEATS = 3           # a picture on this many slides is part of the template
_FIGURES_PER_CALL = 6            # pictures sent in one call: each call carries about 2,000 tokens of fixed overhead
_SMOOTH_EDGE_SHARE = 0.005          # a picture with a smaller share of sharp edges is a soft backdrop

SCENE_INTRO = "Lời mở đầu"
SCENE_OUTRO = "Lời kết"
NOTE_PRESENTER = "Hình ảnh GV"     # not filled in by the AI: the editing notes column is left to the video team

_EMU_PICTURE = 13
_EMU_GROUP = 6


# ── reading the deck ───────────────────────────────────────────────────────────

def _clean(text: Any) -> str:
    return re.sub(r"[ \t ]+", " ", str(text or "")).strip()


def _shape_lines(shape) -> List[str]:
    lines: List[str] = []
    if getattr(shape, "shape_type", None) == _EMU_GROUP:
        for child in shape.shapes:
            lines.extend(_shape_lines(child))
        return lines
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        for paragraph in shape.text_frame.paragraphs:
            text = _clean("".join(run.text for run in paragraph.runs) or paragraph.text)
            if text:
                lines.append(text)
    if getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            cells = [_clean(cell.text) for cell in row.cells]
            if any(cells):
                lines.append(" | ".join(cells))
    if getattr(shape, "has_chart", False) and shape.has_chart:
        text = _chart_text(shape.chart)
        if text:
            lines.append(text)
    return lines


def _chart_text(chart) -> str:
    """A PowerPoint chart as one line of words: its kind, title, and every series with its values."""
    try:
        kind = str(chart.chart_type).split(".")[-1].split(" ")[0].replace("_", " ").lower()
        title = ""
        if chart.has_title and chart.chart_title.has_text_frame:
            title = _clean(chart.chart_title.text_frame.text)
        plot = chart.plots[0]
        categories = [_clean(label) for label in plot.categories]
        parts = []
        for series in list(plot.series)[:6]:
            values = list(series.values)[:24]
            pairs = "; ".join(
                f"{categories[i] if i < len(categories) else i + 1} = {round(v, 4) if isinstance(v, float) else v}"
                for i, v in enumerate(values) if v is not None
            )
            parts.append(f"{_clean(series.name)}: {pairs}" if series.name else pairs)
        head = f"[Biểu đồ {kind}" + (f' "{title}"' if title else "") + "]"
        return (head + " " + " | ".join(parts))[:700]
    except Exception:
        return ""


def _figure_count(shape, slide_area: float) -> int:
    kind = getattr(shape, "shape_type", None)
    if kind == _EMU_GROUP:
        return sum(_figure_count(child, slide_area) for child in shape.shapes)
    area = float(getattr(shape, "width", 0) or 0) * float(getattr(shape, "height", 0) or 0)
    if kind == _EMU_PICTURE:
        return 1 if slide_area and area / slide_area >= _FIGURE_AREA else 0
    if getattr(shape, "has_chart", False) and shape.has_chart:
        return 1
    return 0


def _picture_blobs(shape, slide_area: float) -> List[tuple]:
    """[(share of the slide, image bytes)] of the figure-sized pictures in `shape` (groups included)."""
    kind = getattr(shape, "shape_type", None)
    if kind == _EMU_GROUP:
        found: List[tuple] = []
        for child in shape.shapes:
            found.extend(_picture_blobs(child, slide_area))
        return found
    if kind != _EMU_PICTURE or not slide_area:
        return []
    area = float(getattr(shape, "width", 0) or 0) * float(getattr(shape, "height", 0) or 0)
    share = area / slide_area
    if not _FIGURE_AREA <= share <= _BACKGROUND_AREA:
        return []
    try:
        return [(share, shape.image.blob)]
    except Exception:
        return []


def read_pptx(data: bytes) -> List[Dict[str, Any]]:
    from pptx import Presentation

    deck = Presentation(io.BytesIO(data))
    slide_area = float(deck.slide_width or 1) * float(deck.slide_height or 1)
    slides: List[Dict[str, Any]] = []
    for number, slide in enumerate(deck.slides, start=1):
        title = ""
        try:
            if slide.shapes.title is not None:
                title = _clean(slide.shapes.title.text_frame.text)
        except Exception:
            title = ""
        lines: List[str] = []
        figures = 0
        pictures: List[tuple] = []
        for shape in slide.shapes:
            lines.extend(_shape_lines(shape))
            figures += _figure_count(shape, slide_area)
            pictures.extend(_picture_blobs(shape, slide_area))
        if not title and lines:
            title = lines[0]
        notes = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = _clean(slide.notes_slide.notes_text_frame.text)
        pictures.sort(key=lambda item: item[0], reverse=True)
        slides.append({
            "number": number, "title": title[:200], "text": "\n".join(lines)[:_SLIDE_TEXT_CHARS],
            "notes": notes[:600], "has_figure": figures > 0,
            # The pictures themselves are only carried until their figures have been described.
            "images": [blob for _, blob in pictures[:_MAX_IMAGES_PER_SLIDE]],
        })
    return slides


def read_pdf(data: bytes) -> List[Dict[str, Any]]:
    import pymupdf

    slides: List[Dict[str, Any]] = []
    with pymupdf.open(stream=data, filetype="pdf") as document:
        for number, page in enumerate(document, start=1):
            lines = [_clean(line) for line in page.get_text("text").splitlines()]
            lines = [line for line in lines if line]
            page_area = float(page.rect.width * page.rect.height) or 1.0
            has_figure = False
            pictures = []
            for image in page.get_images(full=True):
                try:
                    for rect in page.get_image_rects(image[0]):
                        share = rect.width * rect.height / page_area
                        if _FIGURE_AREA <= share <= _BACKGROUND_AREA:
                            has_figure = True
                            pictures.append((share, document.extract_image(image[0]).get("image")))
                except Exception:
                    continue
            pictures = sorted((item for item in pictures if item[1]), key=lambda item: item[0], reverse=True)
            slides.append({
                "number": number, "title": (lines[0] if lines else "")[:200],
                "text": "\n".join(lines)[:_SLIDE_TEXT_CHARS], "notes": "", "has_figure": has_figure,
                "images": [blob for _, blob in pictures[:_MAX_IMAGES_PER_SLIDE]],
            })
    return slides


def read_deck(data: bytes, filename: str) -> List[Dict[str, Any]]:
    """[{number, title, text, notes, has_figure}] for every slide of a .pptx or .pdf deck."""
    name = str(filename or "").lower()
    if name.endswith(".pptx"):
        slides = read_pptx(data)
    elif name.endswith(".pdf"):
        slides = read_pdf(data)
    else:
        raise ValueError("Chỉ hỗ trợ file .pptx hoặc .pdf")
    if not slides:
        raise ValueError("Không đọc được slide nào trong file")
    if len(slides) > MAX_SLIDES:
        raise ValueError(f"File có {len(slides)} slide, tối đa {MAX_SLIDES} slide cho một kịch bản")
    # A deck saved as pictures (a scanned or image-only PDF) has nothing to read: writing a script
    # for it would mean inventing the lecture.
    if sum(1 for slide in slides if readable(slide)) < max(1, len(slides) // 3):
        raise ValueError(
            "File này gần như không có chữ đọc được (slide được lưu dưới dạng ảnh). "
            "Hãy dùng file .pptx gốc hoặc PDF xuất trực tiếp từ PowerPoint."
        )
    return slides


def readable(slide: Dict[str, Any]) -> bool:
    """Whether the slide carries enough words (on it or in its notes) to speak about."""
    return len(f"{slide.get('text') or ''} {slide.get('notes') or ''}".split()) >= 4


def has_cover(slides: List[Dict[str, Any]]) -> bool:
    """Whether slide 1 is a title slide: the opening on camera stands in for it."""
    return bool(slides) and len(slides) > 1 and len(str(slides[0].get("text") or "").split()) <= _COVER_WORDS


def looks_vietnamese(slides: List[Dict[str, Any]]) -> bool:
    sample = " ".join(str(slide.get("text") or "") for slide in slides[:12])
    marks = len(re.findall(r"[ăâêôơưđáàảãạấầẩẫậắằẳẵặéèẻẽẹếềểễệíìỉĩịóòỏõọốồổỗộớờởỡợúùủũụứừửữựýỳỷỹỵ]", sample.lower()))
    return marks >= max(3, len(sample) // 200)


# ── looking at the pictures ────────────────────────────────────────────────────

def shrink_image(blob: bytes) -> Optional[str]:
    """A picture as a small JPEG data URL; the model's price grows with the pixels, not the file size."""
    try:
        import base64
        from PIL import Image

        with Image.open(io.BytesIO(blob)) as image:
            image.load()
            if image.mode in ("RGBA", "LA", "P"):
                converted = image.convert("RGBA")
                flat = Image.new("RGB", converted.size, (255, 255, 255))
                flat.paste(converted, mask=converted.split()[-1])
                image = flat
            else:
                image = image.convert("RGB")
            image.thumbnail((_FIGURE_SIDE, _FIGURE_SIDE))
            out = io.BytesIO()
            image.save(out, format="JPEG", quality=82)
        return "data:image/jpeg;base64," + base64.b64encode(out.getvalue()).decode("ascii")
    except Exception as error:
        print(f"[lecture_script] cannot read a picture: {error!r}")
        return None


_DECORATIVE = "TRANG TRI"


def _fold(text: str) -> str:
    import unicodedata
    return "".join(ch for ch in unicodedata.normalize("NFD", str(text or "").upper().replace("Đ", "D")) if unicodedata.category(ch) != "Mn")


def _drop_template_pictures(slides: List[Dict[str, Any]]) -> int:
    """A picture that sits on three or more slides of the deck (a logo, a header band) belongs to the
    template, not to the lesson: it is removed from every slide. Returns how many were removed."""
    import hashlib

    seen: Dict[str, int] = {}
    for slide in slides:
        for digest in {hashlib.sha1(blob).hexdigest() for blob in slide.get("images") or []}:
            seen[digest] = seen.get(digest, 0) + 1
    repeated = {digest for digest, count in seen.items() if count >= _TEMPLATE_REPEATS}
    removed = 0
    for slide in slides:
        kept = [blob for blob in slide.get("images") or [] if hashlib.sha1(blob).hexdigest() not in repeated]
        removed += len(slide.get("images") or []) - len(kept)
        if "images" in slide:
            slide["images"] = kept
    return removed


def is_smooth_picture(blob: bytes) -> bool:
    """True for a picture with no sharp edge at all: a soft gradient or blurred backdrop. There is nothing
    in it to describe, and finding that out costs no model call (in the decks checked, every such
    picture measures exactly 0 on this scale, every diagram, chart and screenshot above 0.02)."""
    try:
        from PIL import Image

        with Image.open(io.BytesIO(blob)) as image:
            image.load()
            gray = image.convert("L")
        gray.thumbnail((160, 160))
        width, height = gray.size
        pixels = gray.load()
        sharp = total = 0
        for y in range(height - 1):
            for x in range(width - 1):
                total += 1
                if max(abs(pixels[x, y] - pixels[x + 1, y]), abs(pixels[x, y] - pixels[x, y + 1])) >= 40:
                    sharp += 1
        return total > 0 and sharp / total < _SMOOTH_EDGE_SHARE
    except Exception:
        return False


async def describe_figures(extractor, slides: List[Dict[str, Any]], vietnamese: bool) -> int:
    """Fills slide["figure"] with what the model sees in the pictures of each slide, then drops the pictures.

    Only content is described. Soft backdrops and pictures repeated across the deck are dropped without
    asking anyone; of the rest, the model itself says TRANG TRÍ for a decoration, in the same call that
    describes its neighbours. A call costs about 2,000 tokens before the first picture, so many pictures
    travel together. A failed call leaves the slides as they were, written from their text alone."""
    cover = has_cover(slides)
    _drop_template_pictures(slides)
    # (slide, picture) pairs worth asking about, in reading order
    pending: List[tuple] = []
    for slide in slides:
        if cover and slide is slides[0]:
            continue
        for blob in slide.get("images") or []:
            if not is_smooth_picture(blob):
                pending.append((slide, blob))
    pending = pending[: _MAX_FIGURE_SLIDES * _MAX_IMAGES_PER_SLIDE]
    gate = asyncio.Semaphore(_PARALLEL_FIGURES)
    model = getattr(extractor, "model_name", None) or None

    async def look(group: List[tuple]) -> List[Optional[str]]:
        urls = [shrink_image(blob) for _, blob in group]
        if not all(urls) or not hasattr(extractor, "_llm_completion_plain_text"):
            return [None] * len(group)
        if vietnamese:
            ask = (
                "Mỗi hình dưới đây nằm trong một slide bài giảng (tiêu đề slide ghi trước hình). Với MỖI hình, viết mô tả để giảng viên giảng cho người nghe không nhìn thấy hình: "
                "loại hình (sơ đồ, biểu đồ, bảng, ảnh chụp màn hình, ảnh minh hoạ), các thành phần chính, nhãn chữ đọc được, mũi tên hay quan hệ giữa chúng. "
                "2-4 câu tiếng Việt, CHỈ tả điều thực sự nhìn thấy: không đoán tên một thành phần nếu hình không ghi, không suy ra từ tiêu đề; chỗ nào không có chữ thì nói là không có chữ. "
                f"Hình chỉ để trang trí (ảnh nền, logo, họa tiết, ảnh người, và dãy biểu tượng nhỏ minh hoạ cho các bước hay ý mà chữ của chúng đã có trên slide, như các icon trong vòng tròn nối nhau) thì mô tả của hình đó là đúng hai chữ: {_DECORATIVE}.\n"
                'Trả về DUY NHẤT JSON: {"items":[{"k":số thứ tự hình,"kind":"diagram|chart|table|screenshot|photo|decoration","has_text":true nếu trong hình có chữ đọc được, "description":"..."}]} '
                "với đúng một phần tử cho mỗi hình. Hình có kind=decoration thì để description là TRANG TRI."
            )
        else:
            ask = (
                "Each picture below sits in a lecture slide (the slide title is written before the picture). For EACH picture write a description for a lecturer who will explain it to "
                "listeners who cannot see it: the kind of picture (diagram, chart, table, screenshot, illustration), its main parts, readable labels, arrows or relations between them. "
                "2-4 sentences, ONLY what is really visible: do not guess the name of a part the picture does not label, do not infer from the title, and say so where there is no text. "
                f"A picture that is only decoration (a background, logo, pattern, a person, and a row of small icons that merely illustrate steps or points whose words are already on the slide, like icons in linked circles) gets exactly the two words: {_DECORATIVE}.\n"
                'Return ONLY JSON: {"items":[{"k":picture number,"kind":"diagram|chart|table|screenshot|photo|decoration","has_text":true if the picture has readable text,"description":"..."}]} '
                "with exactly one item per picture. A picture with kind=decoration gets TRANG TRI as its description."
            )
        content: List[Dict[str, Any]] = [{"type": "text", "text": ask}]
        for number, ((slide, _), url) in enumerate(zip(group, urls), start=1):
            content.append({"type": "text", "text": f"Hình {number} — slide {slide.get('number')}: {slide.get('title') or ''}"})
            content.append({"type": "image_url", "image_url": {"url": url}})
        for _ in range(2):
            async with gate:
                try:
                    raw = await extractor._llm_completion_plain_text(
                        [{"role": "user", "content": content}], max_tokens=300 + 260 * len(group), temperature=0.2,
                        json_mode=True, call_purpose="lecture_script_figure", model_override=model,
                    )
                except Exception as error:
                    print(f"[lecture_script] figures call failed: {error!r}")
                    return [None] * len(group)
            parsed = parse_json_response(raw, clean_result_text=_strip_fence)
            found: Dict[int, str] = {}
            for item in (parsed or {}).get("items") or []:
                try:
                    text = " ".join(str(item.get("description") or "").split())
                    kind = str(item.get("kind") or "").strip().lower()
                    has_text = item.get("has_text")
                    if isinstance(has_text, str):
                        has_text = has_text.strip().lower() in {"true", "yes", "1", "có", "co"}
                    # A diagram or chart with no word in it is a drawing, not something to teach from:
                    # the model's own description can still sound like content ("three blue arrows..."),
                    # so the rule is applied here, not left to its judgement.
                    if kind == "decoration" or (kind in {"diagram", "chart"} and has_text is False):
                        text = _DECORATIVE
                    found[int(item.get("k"))] = text
                except (TypeError, ValueError, AttributeError):
                    continue
            if set(range(1, len(group) + 1)) <= set(found):
                return [found[number] for number in range(1, len(group) + 1)]
        return [None] * len(group)

    groups = [pending[start:start + _FIGURES_PER_CALL] for start in range(0, len(pending), _FIGURES_PER_CALL)]
    try:
        results = await asyncio.gather(*(look(group) for group in groups))
    finally:
        for slide in slides:
            slide.pop("images", None)       # never travels further: not into the task result, not into a saved set
    described = 0
    seen: set = set()
    for group, texts in zip(groups, results):
        for (slide, _), text in zip(group, texts):
            if not text or len(text) < 20 or _fold(text).startswith(_DECORATIVE):
                continue
            slide["figure"] = (str(slide.get("figure") or "") + " " + text).strip()[:700]
            if id(slide) not in seen:
                seen.add(id(slide))
                described += 1
    if pending:
        print(f"[lecture_script] figures: {described} slides described from {len(pending)} pictures in {len(groups)} calls")
    return described


# ── writing the script ─────────────────────────────────────────────────────────

_STYLE_VI = """Bạn viết LỜI THOẠI cho giảng viên đọc khi quay video bài giảng (kịch bản dựng), dựa trên slide có sẵn.

Văn phong bắt buộc:
- Đây là lời NÓI của giảng viên trước sinh viên: tự nhiên, ấm áp, mạch lạc. Mặc định giảng viên xưng "thầy", gọi người học là "các em"; xen kẽ "chúng ta". (Nếu yêu cầu thêm của người dùng nêu cách xưng hô khác thì theo đó.)
- Giảng giải nội dung slide chứ không đọc lại nguyên văn: nêu ý, giải thích vì sao, nối các ý với nhau, thêm ví dụ ngắn khi slide quá cô đọng. Không bịa số liệu, tên riêng, mốc thời gian không có trên slide.
- ĐỘ DÀI: mỗi slide 5-7 đoạn, mỗi đoạn 1-2 câu, mỗi câu không quá 28 từ; tổng cộng khoảng 105-130 từ, KHÔNG BAO GIỜ quá 150 từ (slide rất ít nội dung: 60-90 từ). Giảng giải đủ ý: nói rõ khái niệm, vì sao nó quan trọng, một ví dụ hoặc hình ảnh so sánh ngắn khi slide có nhiều thuật ngữ. Các đoạn cách nhau bằng MỘT DÒNG TRỐNG.
- Mở đầu mỗi slide bằng một câu dẫn tự nhiên, đa dạng, ví dụ: "Chúng ta bắt đầu với...", "Mời các em quan sát...", "Vậy ... là gì?", "Tiếp theo, thầy sẽ...", "Các em thấy đấy,...". Không mở đầu slide nào bằng "Slide này" và không lặp một kiểu mở đầu ở hai slide liền nhau. Câu "Mời các em quan sát..." chỉ dùng cho slide có has_figure=true. Trong các slide được giao, mỗi kiểu mở đầu dùng nhiều nhất một lần.
- Cứ khoảng hai, ba slide thì có một câu hỏi gợi mở cho người học ("Theo các em, ...?", "Vậy ... là gì?") rồi tự trả lời ngay sau đó.
- Slide liệt kê nhiều ý thì đếm thành lời: "thứ nhất..., thứ hai..., và thứ ba...", mỗi ý một đoạn.
- Lời chào đã có ở phần Lời mở đầu: không slide nào chào lại hay "chào mừng" lại người học.
- Kết mỗi slide bằng một câu chốt ý hoặc một câu dẫn sang slide kế tiếp (dựa vào tiêu đề slide kế tiếp được cung cấp).
- Thuật ngữ tiếng Anh giữ nguyên, lần đầu xuất hiện thì đặt trong ngoặc sau từ tiếng Việt, ví dụ "Mối đe dọa (Threat)".
- Công thức, ký hiệu phải viết thành lời để đọc được: "Tấn công bằng Mối đe dọa cộng với Lỗ hổng", "K1 bằng K2".
- Danh sách trên slide được nói thành câu: "thứ nhất là..., thứ hai..., và cuối cùng...".
- Chỉ văn xuôi thuần: không gạch đầu dòng, không markdown, không emoji, không ghi chú trong ngoặc vuông, không viết "(Slide 3)".
- Slide có hình/sơ đồ: nếu có "figure_description" (mô tả hình do AI đã xem) thì giảng giải đúng nội dung hình theo mô tả đó ("Ở sơ đồ này, ..."), nêu các thành phần và quan hệ chính để người nghe hiểu hình mà không cần nhìn; không thêm chi tiết ngoài mô tả. Không có mô tả mà slide ít chữ: mời người học quan sát hình theo tiêu đề, không tả chi tiết bạn không biết.

Ví dụ lời thoại đúng văn phong (chỉ để học giọng văn, không chép nội dung):
---
Chúng ta bắt đầu với phần 2.1: Khái quát về Mối đe dọa, Lỗ hổng và Tấn công.

Thầy muốn hỏi các em một câu: Theo các em, điều gì khiến một hệ thống máy tính không còn an toàn?

Đó chính là sự kết hợp giữa 3 yếu tố: Mối đe dọa (Threat), Điểm yếu (Weakness) và Lỗ hổng (Vulnerability).

Mối đe dọa là hành động gây hại, điểm yếu là lỗi tồn tại trong hệ thống, còn lỗ hổng là điểm yếu cho phép mối đe dọa gây tác hại thực sự.

Các em hãy nhớ rằng, mọi hệ thống đều luôn có điểm yếu, vấn đề là chúng ta quản lý chúng như thế nào.
---
Mời các em quan sát sơ đồ về quan hệ giữa Mối đe dọa và Lỗ hổng.

Các em thấy đấy, mối đe dọa sẽ khai thác lỗ hổng để dẫn đến một cuộc tấn công.

Có ba nguyên lý cốt lõi ở đây: thứ nhất là mối đe dọa thường nhắm vào các lỗ hổng đã biết.

Thứ hai, nếu tồn tại lỗ hổng thì khả năng bị tấn công là rất cao.

Và thứ ba, công thức kinh điển: Tấn công bằng Mối đe dọa cộng với Lỗ hổng.

Nếu chúng ta giảm thiểu được lỗ hổng, khả năng bị tấn công sẽ giảm xuống đáng kể.
---"""

_STYLE_EN = """You write the SPOKEN SCRIPT a lecturer reads when recording a lecture video, based on an existing slide deck.

Required style:
- It is the lecturer speaking to students: natural, warm, clear. Use "I" for the lecturer, "you" for the learners and "we" together (unless the user's extra request asks for something else).
- Explain the slide rather than read it out: state the point, say why it matters, connect the ideas, add a short example when the slide is terse. Never invent figures, names or dates that are not on the slide.
- LENGTH: 5-7 paragraphs per slide, 1-2 sentences each, no sentence over 28 words; about 105-130 words in total and NEVER more than 150 (60-90 for a sparse slide). Explain enough: the concept, why it matters, and a short example or comparison when the slide has many terms. Paragraphs are separated by ONE BLANK LINE.
- Open each slide with a natural, varied lead-in ("Let's start with...", "Take a look at...", "So what is...?"). Never open with "This slide" and never use the same opener on two slides in a row.
- The greeting belongs to the opening scene: no slide greets or welcomes the learners again.
- Close each slide with a takeaway or a bridge to the next slide (its title is given).
- Say formulas and symbols in words. Turn lists into sentences ("first..., second..., and finally...").
- Plain prose only: no bullets, no markdown, no emoji, no bracketed stage directions.
- For a slide with a figure: when "figure_description" is given (what the AI saw in the picture), explain the figure from it ("In this diagram, ..."), naming its parts and relations so a listener understands it without seeing it; add nothing beyond the description. Without a description, on a slide with little text: invite the learners to look at the figure according to the title and do not describe details you do not know."""


def _style(vietnamese: bool) -> str:
    return _STYLE_VI if vietnamese else _STYLE_EN


def _outline(slides: List[Dict[str, Any]]) -> str:
    return "\n".join(f"{slide['number']}. {slide.get('title') or '(không có tiêu đề)'}" for slide in slides)


def _user_request(prompt: str, vietnamese: bool) -> str:
    text = _clean(prompt)[:1500]
    if not text:
        return ""
    head = "Yêu cầu thêm của người dùng (ưu tiên cao hơn mặc định ở trên)" if vietnamese else "Extra request from the user (overrides the defaults above)"
    return f"\n\n{head}:\n{text}"


def _strip_fence(text: str) -> str:
    return str(text or "").strip().removeprefix("```json").removeprefix("```").removesuffix("```").strip()


def tidy_script(text: Any) -> str:
    """Spoken prose only: markdown, bullets and stage directions the model slipped in are removed,
    and paragraphs are separated by exactly one blank line."""
    raw = str(text or "").replace("\r\n", "\n").replace("\r", "\n")
    raw = re.sub(r"\*\*(.+?)\*\*", r"\1", raw)
    raw = re.sub(r"(?m)^\s*(?:[-•*–]\s+|#{1,6}\s+)", "", raw)
    raw = re.sub(r"\[[^\]\n]{0,60}\]", "", raw)
    paragraphs = [" ".join(part.split()) for part in re.split(r"\n\s*\n|\n", raw)]
    return "\n\n".join(part for part in paragraphs if part)


def word_count(text: Any) -> int:
    return len(str(text or "").split())


def estimated_minutes(rows: List[Dict[str, Any]]) -> int:
    words = sum(word_count(row.get("script")) for row in rows)
    return max(1, round(words / WORDS_PER_MINUTE)) if words else 0


async def _ask_json(extractor, messages, *, max_tokens: int, purpose: str) -> Optional[Dict[str, Any]]:
    raw = await extractor._llm_completion_plain_text(
        messages, max_tokens=max_tokens, temperature=0.6, json_mode=True, call_purpose=purpose,
    )
    parsed = parse_json_response(raw, clean_result_text=_strip_fence)
    return parsed if isinstance(parsed, dict) else None


async def _write_batch(
    extractor, deck_title: str, outline: str, batch: List[Dict[str, Any]], next_title: str,
    prompt: str, vietnamese: bool, previous: Optional[Dict[int, str]] = None,
) -> Dict[int, str]:
    payload = {
        "deck_title": deck_title,
        "slides": [
            {
                "slide": slide["number"], "title": slide.get("title") or "", "text_on_slide": slide.get("text") or "",
                "presenter_notes": slide.get("notes") or "", "has_figure": bool(slide.get("has_figure")),
                **({"figure_description": slide["figure"]} if slide.get("figure") else {}),
                **({"current_script": previous[slide["number"]]} if previous and previous.get(slide["number"]) else {}),
            }
            for slide in batch
        ],
        "title_of_the_slide_after_these": next_title,
    }
    task = (
        "Viết lời thoại cho từng slide trong \"slides\". "
        if vietnamese else "Write the script for every slide in \"slides\". "
    )
    if previous:
        task += (
            "Mỗi slide đã có \"current_script\": hãy viết lại theo yêu cầu thêm của người dùng, giữ những gì yêu cầu không đụng tới. "
            if vietnamese else
            "Each slide has a \"current_script\": rewrite it to follow the user's extra request and keep what the request does not touch. "
        )
    system = (
        _style(vietnamese) + _user_request(prompt, vietnamese)
        + ("\n\nDàn ý toàn bộ bài (để nối mạch, không viết lời cho slide ngoài danh sách được giao):\n" if vietnamese
           else "\n\nOutline of the whole deck (for continuity only; write nothing for slides you were not given):\n")
        + outline
        + "\n\n" + task
        + ("Trả về DUY NHẤT JSON: {\"items\":[{\"slide\":số,\"script\":\"lời thoại, các đoạn cách nhau bằng \\n\\n\"}]} "
           "với đúng một phần tử cho mỗi slide được giao, theo đúng thứ tự."
           if vietnamese else
           "Return ONLY JSON: {\"items\":[{\"slide\":number,\"script\":\"the script, paragraphs separated by \\n\\n\"}]} "
           "with exactly one item per slide given, in order.")
    )
    messages = [{"role": "system", "content": system}, {"role": "user", "content": json.dumps(payload, ensure_ascii=False)}]
    wanted = {slide["number"] for slide in batch}
    found: Dict[int, str] = {}
    for _ in range(2):
        try:
            parsed = await _ask_json(extractor, messages, max_tokens=700 + 520 * len(batch), purpose="lecture_script_slides")
        except Exception as error:
            print(f"[lecture_script] batch {sorted(wanted)} failed: {error!r}")
            parsed = None
        for item in (parsed or {}).get("items") or []:
            if not isinstance(item, dict):
                continue
            try:
                number = int(item.get("slide"))
            except (TypeError, ValueError):
                continue
            script = tidy_script(item.get("script"))
            if number in wanted and word_count(script) >= 12:
                found.setdefault(number, script)
        if wanted <= set(found):
            break
    return found


async def _write_frame(
    extractor, deck_title: str, outline: str, prompt: str, vietnamese: bool,
    previous: Optional[Dict[str, str]] = None, full_intro: bool = False, cover_text: str = "",
) -> Dict[str, str]:
    """The two on-camera scenes: the opening before slide 2 and the closing after the last slide.

    Only the first video of a course opens by welcoming the learners and introducing the lecturer and
    the course; every other video opens with a short greeting and what this part is about."""
    if vietnamese:
        if full_intro:
            intro_spec = (
                "- \"intro\" (Lời mở đầu, 70-110 từ, 4-6 đoạn ngắn cách nhau bằng dòng trống): chào mừng người học, "
                "giới thiệu giảng viên và học phần/môn học (CHỈ những gì có ghi trên slide bìa trong \"cover_slide_text\"), "
                "nói ngắn gọn bài này sẽ giúp họ hiểu điều gì, mời vào nội dung. Kiểu: \"Chào các em.\\n\\nChào mừng các em đến với môn học...\".\n"
            )
        else:
            intro_spec = (
                "- \"intro\" (Lời mở đầu, 30-55 từ, 2-3 đoạn ngắn cách nhau bằng dòng trống): chỉ chào ngắn rồi nói phần này học gì "
                "(theo tiêu đề và dàn ý), mời vào nội dung. TUYỆT ĐỐI không chào mừng dài dòng, không giới thiệu tên giảng viên, "
                "tên trường hay tên học phần (đã giới thiệu ở video đầu tiên). "
                "Kiểu: \"Chào các em.\\n\\nTrong video này, chúng ta sẽ tìm hiểu về...\\n\\nCác em hãy cùng thầy đi vào nội dung chi tiết nhé.\".\n"
            )
        task = (
            "Viết hai đoạn giảng viên nói trước ống kính (không có slide):\n"
            + intro_spec
            + "- \"outro\" (Lời kết, 30-55 từ, MỘT đoạn): tóm lại vừa học xong phần gì, hẹn nội dung tiếp theo, chào. "
            "Kiểu: \"Vậy là chúng ta đã hoàn thành phần ... Xin chào và hẹn gặp lại các em.\"\n"
            "Không bịa tên giảng viên, tên trường hay tên môn học nếu slide không ghi."
        )
        task += (
            "\n- \"title\": tên ngắn của bài giảng này (tối đa 14 từ) nêu đúng các nội dung chính trong dàn ý, "
            "ví dụ \"Mối đe dọa, Lỗ hổng và Phân loại tấn công\"; không ghi tên trường, tên giảng viên."
        )
        shape = "Trả về DUY NHẤT JSON: {\"title\":\"...\",\"intro\":\"...\",\"outro\":\"...\"}"
    else:
        if full_intro:
            intro_spec = (
                "- \"intro\" (70-110 words, 4-6 short paragraphs separated by a blank line): welcome the learners, introduce the lecturer "
                "and the course (ONLY what the title slide in \"cover_slide_text\" says), say briefly what they will understand after it, "
                "lead into the content.\n"
            )
        else:
            intro_spec = (
                "- \"intro\" (30-55 words, 2-3 short paragraphs separated by a blank line): a short greeting, then what this part covers "
                "(from the title and outline), then lead into the content. Do NOT give a long welcome and do NOT introduce the lecturer, "
                "school or course (the first video did that).\n"
            )
        task = (
            "Write the two scenes the lecturer speaks on camera (no slide):\n"
            + intro_spec
            + "- \"outro\" (30-55 words, ONE paragraph): sum up what was covered, point to what comes next, say goodbye.\n"
            "Do not invent a lecturer, school or course name that the slides do not give."
        )
        task += "\n- \"title\": a short name for this lecture (at most 14 words) built from the main topics of the outline."
        shape = "Return ONLY JSON: {\"title\":\"...\",\"intro\":\"...\",\"outro\":\"...\"}"
    payload: Dict[str, Any] = {"deck_title": deck_title, "outline": outline}
    if full_intro and cover_text:
        payload["cover_slide_text"] = cover_text[:600]
    if previous:
        payload["current"] = previous
        task += ("\nĐã có bản hiện tại trong \"current\": viết lại theo yêu cầu thêm của người dùng." if vietnamese
                 else "\nThe current version is in \"current\": rewrite it to follow the user's extra request.")
    system = _style(vietnamese) + _user_request(prompt, vietnamese) + "\n\n" + task + "\n" + shape
    messages = [{"role": "system", "content": system}, {"role": "user", "content": json.dumps(payload, ensure_ascii=False)}]
    for _ in range(2):
        try:
            parsed = await _ask_json(extractor, messages, max_tokens=900, purpose="lecture_script_frame")
        except Exception as error:
            print(f"[lecture_script] intro/outro failed: {error!r}")
            parsed = None
        intro, outro = tidy_script((parsed or {}).get("intro")), tidy_script((parsed or {}).get("outro"))
        if word_count(intro) >= 10 and word_count(outro) >= 8:
            title = " ".join(str((parsed or {}).get("title") or "").split()).strip(" .\"'")
            return {"intro": intro, "outro": outro.replace("\n\n", " "), "title": title if 2 <= len(title) <= 160 else ""}
    return {}


def _fallback_frame(deck_title: str, vietnamese: bool) -> Dict[str, str]:
    topic = deck_title or ("bài học hôm nay" if vietnamese else "today's lecture")
    if vietnamese:
        return {
            "intro": f"Chào các em.\n\nHôm nay, thầy và các em sẽ cùng tìm hiểu nội dung: {topic}.\n\nCác em hãy cùng thầy đi vào nội dung chi tiết nhé.",
            "outro": f"Vậy là chúng ta đã hoàn thành phần {topic}. Xin chào và hẹn gặp lại các em.",
        }
    return {
        "intro": f"Hello everyone.\n\nToday we will work through: {topic}.\n\nLet's get started.",
        "outro": f"That brings us to the end of {topic}. Thank you, and see you next time.",
    }


def deck_title_of(slides: List[Dict[str, Any]], filename: str = "") -> str:
    first = slides[0] if slides else {}
    lines = [line for line in str(first.get("text") or "").splitlines() if line.strip()]
    if has_cover(slides) and lines:
        # A title slide often stacks school, course and chapter: the last two lines name the lecture.
        named = [line for line in lines if 3 <= len(line) <= 90 and not re.search(r"ths\.|ts\.|pgs\.|gv\b|@|khoa\s", line.lower())]
        if named:
            return " – ".join(named[-2:]) if len(named) > 2 else " – ".join(named)
    title = _clean(first.get("title"))
    if title:
        return title
    return re.sub(r"\.(pptx|pdf)$", "", str(filename or ""), flags=re.IGNORECASE)


def assemble_rows(slides: List[Dict[str, Any]], scripts: Dict[int, str], frame: Dict[str, str]) -> List[Dict[str, Any]]:
    """Rows in the order the video is cut: opening, one row per slide, closing."""
    rows: List[Dict[str, Any]] = [
        {"scene": SCENE_INTRO, "slide": 1 if has_cover(slides) else None, "script": frame.get("intro", ""), "note": ""}
    ]
    for slide in slides[1:] if has_cover(slides) else slides:
        number = slide["number"]
        rows.append({
            "scene": f"Slide {number}", "slide": number, "script": scripts.get(number, ""),
            "note": "",
        })
    rows.append({"scene": SCENE_OUTRO, "slide": None, "script": frame.get("outro", ""), "note": ""})
    return rows


Progress = Callable[[int], Awaitable[None]]


async def write_script(
    extractor, slides: List[Dict[str, Any]], *, prompt: str = "", filename: str = "",
    previous_rows: Optional[List[Dict[str, Any]]] = None, on_progress: Optional[Progress] = None,
    first_video: bool = False, look_at_figures: bool = True,
) -> Dict[str, Any]:
    """The whole script for `slides`. With `previous_rows`, an existing script is rewritten to
    follow `prompt` instead of being written from nothing."""
    vietnamese = looks_vietnamese(slides)
    if look_at_figures:
        figures = await describe_figures(extractor, slides, vietnamese)
    else:
        figures = 0
        for slide in slides:
            slide.pop("images", None)
    deck_title = deck_title_of(slides, filename)
    outline = _outline(slides)
    body = slides[1:] if has_cover(slides) else slides

    previous_scripts: Dict[int, str] = {}
    previous_frame: Dict[str, str] = {}
    for row in previous_rows or []:
        if not isinstance(row, dict) or not str(row.get("script") or "").strip():
            continue
        if row.get("scene") == SCENE_INTRO:
            previous_frame["intro"] = str(row["script"])
        elif row.get("scene") == SCENE_OUTRO:
            previous_frame["outro"] = str(row["script"])
        elif isinstance(row.get("slide"), int):
            previous_scripts[row["slide"]] = str(row["script"])

    # A slide with no words is left for the user to write: nothing is made up for it.
    body = [slide for slide in body if readable(slide) or previous_scripts.get(slide["number"])]
    batches = [body[start:start + _BATCH_SLIDES] for start in range(0, len(body), _BATCH_SLIDES)]
    done = 0
    gate = asyncio.Semaphore(_PARALLEL_BATCHES)

    async def run(index: int, batch: List[Dict[str, Any]]) -> Dict[int, str]:
        nonlocal done
        following = batches[index + 1][0].get("title", "") if index + 1 < len(batches) else ""
        async with gate:
            result = await _write_batch(
                extractor, deck_title, outline, batch, following, prompt, vietnamese, previous_scripts or None,
            )
        done += 1
        if on_progress:
            await on_progress(10 + int(80 * done / max(1, len(batches) + 1)))
        return result

    frame_task = asyncio.create_task(
        _write_frame(
            extractor, deck_title, outline, prompt, vietnamese, previous_frame or None,
            full_intro=first_video, cover_text=str(slides[0].get("text") or "") if has_cover(slides) else "",
        )
    )
    results = await asyncio.gather(*(run(index, batch) for index, batch in enumerate(batches)))
    scripts: Dict[int, str] = {}
    for result in results:
        scripts.update(result)
    frame = await frame_task or {}
    if previous_frame:
        frame = {**previous_frame, **frame}
    frame = {**_fallback_frame(deck_title, vietnamese), **{key: value for key, value in frame.items() if value}}
    # A slide the model skipped keeps what it had before; a slide never written stays empty for the user to fill.
    for number, text in previous_scripts.items():
        scripts.setdefault(number, text)

    rows = assemble_rows(slides, scripts, frame)
    missing = [row["scene"] for row in rows if not row["script"].strip()]
    if len(missing) > max(1, len(rows) // 2):
        raise RuntimeError("AI chưa viết được kịch bản cho file này. Hãy thử lại sau ít phút.")
    return {
        "title": frame.get("title") or deck_title,
        "language": "vi" if vietnamese else "en",
        "slide_count": len(slides),
        "figures_described": figures,
        "duration_minutes": estimated_minutes(rows),
        "word_count": sum(word_count(row["script"]) for row in rows),
        "rows": rows,
        "missing": missing,
    }


# ── English subtitles ──────────────────────────────────────────────────────────

_TRANSLATE_BATCH_CHARS = 2600
TRANSLATION_TARGETS = ("en", "vi")

_TO_ENGLISH = """You turn the spoken script of a Vietnamese university lecturer into English subtitles for a lecture video.

Translate the MEANING and write what a good English-speaking lecturer would actually say to students. Never translate word by word:
- Build each sentence the English way. Inside one paragraph you may reorder, split or join sentences when it reads better.
- Drop Vietnamese filler and politeness particles (nhé, nhỉ, đấy, thế, ạ, "các em thấy đấy" said out of habit) unless they carry meaning.
- No calques: "khiến cho hệ thống gặp nguy hiểm" is "puts a system at risk", not "makes the system meet danger".
- Spoken and plain, contractions welcome, no stiff written phrases ("in order to", "it is necessary that", "regarding").
- Subtitles are read fast: keep sentences short, about 20 words at most.
- Who is speaking: "thầy / cô / tôi" is "I", "các em / các bạn" is "you", "chúng ta / thầy trò mình" is "we".
- Keep technical terms, acronyms, product names and every number exactly. Where the Vietnamese gives a term followed by its English original in parentheses, use the English term once and drop the parentheses.
- Natural does not mean shorter: keep every piece of meaning (who hopes or asks, where something is, each item of a list). Keep every fact and add none.

Examples of the difference:
VI: Mời các em quan sát sơ đồ về quan hệ giữa Mối đe dọa và Lỗ hổng.
Word by word (wrong): Please observe the diagram about the relationship between Threat and Vulnerability.
Natural (right): Take a look at this diagram of how threats and vulnerabilities relate.
VI: Các em thấy đấy, mối đe dọa sẽ khai thác lỗ hổng để dẫn đến một cuộc tấn công.
Word by word (wrong): You see, the threat will exploit the vulnerability to lead to an attack.
Natural (right): A threat exploits a vulnerability, and that's what leads to an attack.
VI: Vì vậy, loại bỏ lỗ hổng là cách hiệu quả nhất để ngăn chặn tấn công thành công.
Natural (right): So the most effective way to stop an attack is to get rid of the vulnerability."""

_TO_VIETNAMESE = """Bạn chuyển lời thoại tiếng Anh của giảng viên thành phụ đề tiếng Việt cho video bài giảng đại học.

Dịch Ý, viết đúng như một giảng viên người Việt sẽ nói với sinh viên. Tuyệt đối không dịch từng chữ:
- Đặt câu theo cách nói tiếng Việt. Trong một đoạn được đảo trật tự, tách hoặc gộp câu cho tự nhiên.
- Không bê nguyên cấu trúc tiếng Anh: tránh "nó là ... mà", tránh lạm dụng "bị / được", "một cách", "của" chồng nhau.
- Văn nói, gần gũi, rõ ý; câu ngắn để đọc kịp phụ đề.
- Xưng hô: "I" là "thầy" (trừ khi yêu cầu thêm nêu cách xưng khác), "you" là "các em", "we" là "chúng ta".
- Thuật ngữ chuyên ngành, từ viết tắt, tên riêng và mọi con số giữ nguyên. Thuật ngữ đã quen bằng tiếng Việt thì dùng tiếng Việt và để từ tiếng Anh trong ngoặc ở lần đầu, ví dụ "Mối đe dọa (Threat)".
- Tự nhiên không có nghĩa là lược bớt: giữ đủ mọi ý (ai mong, ai hỏi, ở đâu, từng mục trong danh sách), không thêm ý.

Ví dụ:
EN: Take a look at this diagram of how threats and vulnerabilities relate.
Dịch từng chữ (sai): Hãy nhìn vào sơ đồ này của cách các mối đe dọa và lỗ hổng liên hệ.
Tự nhiên (đúng): Mời các em quan sát sơ đồ về quan hệ giữa mối đe dọa và lỗ hổng.
EN: So the most effective way to stop an attack is to get rid of the vulnerability.
Tự nhiên (đúng): Vì vậy, loại bỏ lỗ hổng là cách hiệu quả nhất để ngăn một cuộc tấn công."""

_TRANSLATE_CONTRACT = (
    "\n\nYou get JSON {\"lecture\": title, \"rows\":[{\"i\": number, \"paragraphs\":[...]}]}. Each row is the script of one scene; "
    "its paragraphs are the subtitle lines, so they must stay aligned: one paragraph in gives exactly one paragraph out, in the "
    "same order, never merged, never dropped. Use the other paragraphs of the row as context.\n"
    "Return ONLY JSON {\"rows\":[{\"i\": number, \"paragraphs\":[...]}]} with the same rows and the same number of paragraphs in each."
)


def _paragraphs(script: Any) -> List[str]:
    return [part.strip() for part in re.split(r"\n\s*\n|\n", str(script or "")) if part.strip()]


async def _translate_rows(
    extractor, rows: List[tuple], target: str, title: str, note: str,
) -> Dict[int, List[str]]:
    """{row index: translated paragraphs} for the rows the model returned with the right paragraph count."""
    system = (_TO_ENGLISH if target == "en" else _TO_VIETNAMESE) + _TRANSLATE_CONTRACT
    if _clean(note):
        system += "\n\nExtra instruction from the user (about how the lecturer speaks; follow it): " + _clean(note)[:600]
    payload = {"lecture": title, "rows": [{"i": index, "paragraphs": paragraphs} for index, paragraphs in rows]}
    messages = [{"role": "system", "content": system}, {"role": "user", "content": json.dumps(payload, ensure_ascii=False)}]
    size = sum(len(text) for _, paragraphs in rows for text in paragraphs)
    try:
        parsed = await _ask_json(extractor, messages, max_tokens=min(7000, 500 + size), purpose="lecture_script_translate")
    except Exception as error:
        print(f"[lecture_script] translation call failed: {error!r}")
        return {}
    wanted = {index: len(paragraphs) for index, paragraphs in rows}
    found: Dict[int, List[str]] = {}
    for item in (parsed or {}).get("rows") or []:
        try:
            index = int(item.get("i"))
            paragraphs = [" ".join(str(text or "").split()) for text in (item.get("paragraphs") or [])]
        except (TypeError, ValueError, AttributeError):
            continue
        if wanted.get(index) == len(paragraphs) and all(paragraphs):
            found[index] = paragraphs
    return found


async def translate_scripts(
    extractor, scripts: List[str], target: str = "en", *, title: str = "", note: str = "",
    on_progress: Optional[Progress] = None,
) -> Dict[str, Any]:
    """{"items": [{"i": row index, "text": "..."}], "failed": [row index]}: `scripts` in the `target` language.

    A row (the script of one scene) is translated as a whole, so the sentences are written the way the
    target language says them instead of word by word, but paragraph for paragraph: the lines of the
    two languages stay side by side. A row the model returns with the wrong number of paragraphs is
    asked again on its own; one that still fails is reported, never half-translated."""
    target = target if target in TRANSLATION_TARGETS else "en"
    rows = [(index, _paragraphs(script)) for index, script in enumerate(scripts)]
    rows = [row for row in rows if row[1]]
    batches: List[List[tuple]] = []
    current: List[tuple] = []
    size = 0
    for row in rows:
        length = sum(len(text) for text in row[1])
        if current and size + length > _TRANSLATE_BATCH_CHARS:
            batches.append(current)
            current, size = [], 0
        current.append(row)
        size += length
    if current:
        batches.append(current)

    done = 0
    gate = asyncio.Semaphore(_PARALLEL_BATCHES)

    async def run(batch: List[tuple]) -> Dict[int, List[str]]:
        nonlocal done
        async with gate:
            found = await _translate_rows(extractor, batch, target, title, note)
            for row in batch:                                   # a row that came back wrong gets one more try, alone
                if row[0] not in found:
                    found.update(await _translate_rows(extractor, [row], target, title, note))
        done += 1
        if on_progress:
            await on_progress(5 + int(90 * done / max(1, len(batches))))
        return found

    translated: Dict[int, List[str]] = {}
    for found in await asyncio.gather(*(run(batch) for batch in batches)):
        translated.update(found)
    items = [{"i": index, "text": "\n\n".join(tidy_script(text) for text in translated[index])} for index, _ in rows if index in translated]
    failed = [index for index, _ in rows if index not in translated]
    return {"items": items, "failed": failed, "target": target}
